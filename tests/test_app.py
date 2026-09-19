from io import BytesIO
from zipfile import ZipFile

from docx import Document

import app


def _template_bytes(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def test_generate_document_replaces_braced_placeholders():
    result = app.generate_document(
        _template_bytes("Patient: {{patient_name}} | Vaccine: {{vaccine}}"),
        {"patient_name": "Alex Example", "vaccine": "Tdap"},
    )
    doc = Document(BytesIO(result))
    assert "Patient: Alex Example | Vaccine: Tdap" in "\n".join(p.text for p in doc.paragraphs)


def test_generate_document_replaces_placeholders_in_tables():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "{{patient_name}}"
    table.cell(0, 1).text = "{{vaccine}}"
    raw = BytesIO()
    doc.save(raw)

    result = app.generate_document(
        raw.getvalue(),
        {"patient_name": "Example Person", "vaccine": "MMR"},
    )
    rendered = Document(BytesIO(result))
    assert rendered.tables[0].cell(0, 0).text == "Example Person"
    assert rendered.tables[0].cell(0, 1).text == "MMR"


def test_make_zip_creates_one_docx_per_record_and_manifest():
    records = [
        {"patient_name": "Alpha", "vaccine": "MMR"},
        {"patient_name": "Beta", "vaccine": "Hepatitis B"},
    ]
    zip_bytes, manifest = app.make_zip(
        _template_bytes("{{patient_name}} - {{vaccine}}"), records
    )
    with ZipFile(BytesIO(zip_bytes)) as archive:
        docx_names = [name for name in archive.namelist() if name.endswith(".docx")]
        assert "manifest.csv" in archive.namelist()
    assert len(docx_names) == 2
    assert len(manifest) == 2


def test_database_configured_requires_url_and_driver(monkeypatch):
    monkeypatch.delenv("DATABASE_URL", raising=False)
    monkeypatch.setattr(app, "psycopg", object(), raising=False)
    assert app.database_configured() is False
    monkeypatch.setenv("DATABASE_URL", "postgresql://example")
    assert app.database_configured() is True


def test_saved_template_bytes_are_reusable():
    template = app.VaccineTemplate(
        id="t1", name="Travel Vaccine Rx", description="test", filename="rx.docx", bytes_data=b"docx-bytes"
    )
    assert app.get_template_bytes(template) == b"docx-bytes"
