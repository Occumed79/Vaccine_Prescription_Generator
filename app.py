from __future__ import annotations

import hashlib
import io
import os
import re
import uuid
import zipfile
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Mapping, Sequence

from docx import Document

try:
    import psycopg
    from psycopg.rows import dict_row
except Exception:
    psycopg = None
    dict_row = None

TOKEN_RE = re.compile(r"\{\{\s*([A-Za-z0-9_.-]+)\s*\}\}")

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass
class VaccineTemplate:
    id: str
    name: str
    description: str
    filename: str
    bytes_data: bytes
    file_sha256: str = ""


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def get_database_url() -> str:
    return os.getenv("DATABASE_URL", "").strip()


def database_configured() -> bool:
    return bool(get_database_url()) and psycopg is not None


def get_db_connection():
    if not database_configured():
        raise RuntimeError("DATABASE_URL is not configured or psycopg is not installed.")
    return psycopg.connect(get_database_url(), row_factory=dict_row, autocommit=True)


def ensure_template_table() -> None:
    with get_db_connection() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS vaccine_prescription_templates (
                id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                description TEXT NOT NULL DEFAULT '',
                filename TEXT NOT NULL,
                content_type TEXT NOT NULL DEFAULT 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                file_bytes BYTEA NOT NULL,
                file_sha256 TEXT NOT NULL,
                is_active BOOLEAN NOT NULL DEFAULT TRUE,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            );
            """
        )


def save_template_to_neon(name: str, description: str, filename: str, template_bytes: bytes) -> str:
    ensure_template_table()
    template_id = f"vaccine-{uuid.uuid4().hex}"
    with get_db_connection() as conn:
        conn.execute(
            """
            INSERT INTO vaccine_prescription_templates
                (id, name, description, filename, content_type, file_bytes, file_sha256, is_active)
            VALUES
                (%(id)s, %(name)s, %(description)s, %(filename)s, %(content_type)s, %(file_bytes)s, %(file_sha256)s, TRUE);
            """,
            {
                "id": template_id,
                "name": name,
                "description": description,
                "filename": filename,
                "content_type": DOCX_MIME,
                "file_bytes": template_bytes,
                "file_sha256": sha256_bytes(template_bytes),
            },
        )
    return template_id


def load_templates_from_neon() -> list[VaccineTemplate]:
    ensure_template_table()
    with get_db_connection() as conn:
        rows = conn.execute(
            """
            SELECT id, name, description, filename, file_bytes, file_sha256
            FROM vaccine_prescription_templates
            WHERE is_active = TRUE
            ORDER BY name ASC;
            """
        ).fetchall()
    return [
        VaccineTemplate(
            id=row["id"],
            name=row["name"],
            description=row["description"] or "",
            filename=row["filename"],
            bytes_data=bytes(row["file_bytes"]),
            file_sha256=row["file_sha256"] or "",
        )
        for row in rows
    ]


def get_all_templates() -> tuple[list[VaccineTemplate], str]:
    if database_configured():
        try:
            return load_templates_from_neon(), "Neon prescription-form library connected."
        except Exception as exc:
            return [], f"Neon is configured but unavailable: {exc}"
    return [], "Neon is not configured. Add DATABASE_URL to use the saved prescription-form library."


def get_template_bytes(template: VaccineTemplate) -> bytes:
    return template.bytes_data



def _string_values(values: Mapping[str, object]) -> dict[str, str]:
    return {str(key): "" if value is None else str(value) for key, value in values.items()}


def _render_text(text: str, values: Mapping[str, str]) -> str:
    return TOKEN_RE.sub(lambda match: values.get(match.group(1), match.group(0)), text)


def _replace_in_paragraph(paragraph, values: Mapping[str, str]) -> int:
    original = paragraph.text
    rendered = _render_text(original, values)
    if rendered == original:
        return 0
    if paragraph.runs:
        paragraph.runs[0].text = rendered
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(rendered)
    return 1


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def replace_placeholders(document: Document, values: Mapping[str, object]) -> int:
    """Replace ``{{column_name}}`` tokens in body paragraphs and tables."""
    normalized = _string_values(values)
    count = 0
    for paragraph in document.paragraphs:
        count += _replace_in_paragraph(paragraph, normalized)
    for table in document.tables:
        for paragraph in _iter_table_paragraphs(table):
            count += _replace_in_paragraph(paragraph, normalized)
    return count


def generate_document(template_bytes: bytes, values: Mapping[str, object]) -> bytes:
    document = Document(io.BytesIO(template_bytes))
    replace_placeholders(document, values)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _safe_filename(value: object, fallback: str) -> str:
    text = str(value or "").strip()
    text = re.sub(r"[^A-Za-z0-9._ -]+", "", text).strip().replace(" ", "-")
    return text[:80] or fallback


def make_zip(
    template_bytes: bytes,
    records: Sequence[Mapping[str, object]],
) -> tuple[bytes, list[dict[str, str]]]:
    """Generate one prescription document per record and package them in a ZIP."""
    archive_buffer = io.BytesIO()
    manifest: list[dict[str, str]] = []

    with zipfile.ZipFile(archive_buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for index, record in enumerate(records, start=1):
            rendered = generate_document(template_bytes, record)
            identity = (
                record.get("patient_name")
                or record.get("examinee_name")
                or record.get("name")
                or f"prescription-{index}"
            )
            stem = _safe_filename(identity, f"prescription-{index}")
            filename = f"{index:03d}-{stem}.docx"
            archive.writestr(filename, rendered)
            manifest.append({"row": str(index), "file": filename})

        manifest_text = "row,file\n" + "".join(
            f'{item["row"]},{item["file"]}\n' for item in manifest
        )
        archive.writestr("manifest.csv", manifest_text.encode("utf-8"))

    return archive_buffer.getvalue(), manifest


def _records_from_dataframe(frame) -> list[dict[str, object]]:
    clean = frame.where(frame.notna(), "")
    return clean.to_dict(orient="records")


def _read_source_file(source):
    import pandas as pd

    suffix = Path(source.name).suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(source, dtype=str).fillna("")
    return pd.read_excel(source, dtype=str).fillna("")


def render_generate_tab() -> None:
    import streamlit as st

    templates, status = get_all_templates()
    st.caption(status)

    selected_template_bytes = None
    if templates:
        labels = ["Upload a one-time prescription form"] + [template.name for template in templates]
        selected = st.selectbox("Prescription form", labels)
        if selected == labels[0]:
            uploaded_template = st.file_uploader("Vaccine prescription Word template", type=["docx"], key="one_time_template")
            if uploaded_template is not None:
                selected_template_bytes = uploaded_template.getvalue()
        else:
            template = next(item for item in templates if item.name == selected)
            selected_template_bytes = get_template_bytes(template)
            if template.description:
                st.caption(template.description)
    else:
        uploaded_template = st.file_uploader("Vaccine prescription Word template", type=["docx"], key="fallback_template")
        if uploaded_template is not None:
            selected_template_bytes = uploaded_template.getvalue()

    source = st.file_uploader("Prescription data", type=["xlsx", "xlsm", "xls", "csv"])
    if selected_template_bytes is None or source is None:
        st.info("Choose a saved prescription form (or upload one) and add prescription data to continue.")
        return

    try:
        frame = _read_source_file(source)
    except Exception as exc:
        st.error(f"Could not read prescription data: {exc}")
        return

    if frame.empty:
        st.warning("The uploaded spreadsheet has no prescription rows.")
        return

    st.subheader("Prescription Preview")
    st.dataframe(frame, use_container_width=True, hide_index=True)
    st.caption("Template tokens use spreadsheet column names, for example {{patient_name}} or {{vaccine}}.")

    if st.button("Generate Vaccine Prescriptions", type="primary"):
        try:
            zip_bytes, manifest = make_zip(selected_template_bytes, _records_from_dataframe(frame))
            st.success(f"Generated {len(manifest)} vaccine prescription document(s).")
            st.download_button(
                "Download ZIP",
                data=zip_bytes,
                file_name=f"vaccine-prescriptions-{date.today().isoformat()}.zip",
                mime="application/zip",
            )
        except Exception as exc:
            st.error(f"Prescription generation failed: {exc}")


def render_configure_tab() -> None:
    import streamlit as st

    st.subheader("Saved Prescription Forms")
    templates, status = get_all_templates()
    st.caption(status)
    if templates:
        st.dataframe(
            [{"Name": t.name, "Description": t.description, "Filename": t.filename} for t in templates],
            use_container_width=True,
            hide_index=True,
        )

    st.subheader("Add Prescription Form")
    with st.form("save_vaccine_template", clear_on_submit=True):
        name = st.text_input("Form name")
        description = st.text_area("Description")
        uploaded = st.file_uploader("Prescription form (.docx)", type=["docx"])
        submitted = st.form_submit_button("Save form to Neon")

    if submitted:
        if not database_configured():
            st.error("Neon is not connected. Add DATABASE_URL in Render before saving forms.")
            return
        if not name.strip() or uploaded is None:
            st.error("Form name and .docx file are required.")
            return
        try:
            save_template_to_neon(name.strip(), description.strip(), uploaded.name, uploaded.getvalue())
            st.success("Prescription form saved permanently in Neon.")
            st.rerun()
        except Exception as exc:
            st.error(f"Could not save prescription form: {exc}")


def render_cdc_data_tab() -> None:
    import pandas as pd
    import streamlit as st
    import streamlit.components.v1 as components

    from cdc_data import (
        CDC_SODA2_SOURCES,
        fetch_soda2_rows,
        get_socrata_app_token,
        source_categories,
        sources_for_category,
    )
    from cdc_reference import (
        CDC_CLINICAL_SOURCES,
        CDC_CONTENT_SOURCES,
        CDC_IIS_TABLE_SOURCES,
        CDSI_DOWNLOADS,
        IIS_ACCESS_OPTIONS_URL,
        IIS_RUNTIME_REST_URL,
        PARKED_IMPORT_SOURCES,
        TRAVEL_HEALTH_NOTICES_RSS,
        TRAVEL_HEALTH_NOTICES_URL,
        YELLOW_BOOK_COUNTRY_URL,
        fetch_content_api_html,
        fetch_html_tables,
        fetch_travel_destinations,
        fetch_travel_health_notices,
        fetch_travel_vaccine_table,
    )

    st.subheader("CDC Adult Vaccine Data")
    st.caption(
        "Live CDC data, official adult schedules, vaccine code sets, and destination-specific "
        "travel vaccine recommendations in one workspace."
    )

    live_tab, schedule_tab, clinical_tab, codes_tab, travel_tab, notices_tab, imports_tab = st.tabs(
        [
            "SODA2 Data",
            "Adult Schedules",
            "Clinical Rules",
            "Vaccine Codes",
            "Travel Vaccines",
            "Travel Notices",
            "Importer Queue",
        ]
    )

    with live_tab:
        st.caption(
            f"{len(CDC_SODA2_SOURCES)} live CDC SODA2 sources are registered. "
            + (
                "Socrata App Token connected."
                if get_socrata_app_token()
                else "No Socrata App Token detected; public anonymous access will be attempted."
            )
        )

        category = st.selectbox(
            "Data category",
            ["All"] + source_categories(),
            key="cdc_category",
        )
        sources = sources_for_category(category)
        selected_label = st.selectbox(
            "CDC dataset",
            [f"{source.title} · {source.dataset_id}" for source in sources],
            key="cdc_source",
        )
        source = next(item for item in sources if selected_label.endswith(item.dataset_id))

        st.markdown(f"**{source.title}**")
        st.caption(source.description)
        st.caption(f"Audience: {source.audience} · Dataset ID: {source.dataset_id}")

        col1, col2 = st.columns([1, 2])
        with col1:
            row_limit = st.select_slider(
                "Preview rows",
                options=[25, 50, 100, 250, 500, 1000],
                value=100,
                key="cdc_limit",
            )
        with col2:
            where_clause = st.text_input(
                "Optional SoQL filter",
                placeholder="Example: year = '2025'",
                key="cdc_where",
            )

        if st.button("Load CDC data", type="primary", key="load_cdc_data"):
            try:
                rows = fetch_soda2_rows(
                    source.dataset_id,
                    limit=int(row_limit),
                    where=where_clause.strip() or None,
                )
                st.session_state["cdc_rows"] = rows
                st.session_state["cdc_loaded_id"] = source.dataset_id
                st.session_state["cdc_loaded_title"] = source.title
            except Exception as exc:
                st.session_state.pop("cdc_rows", None)
                st.error(f"CDC SODA2 request failed: {exc}")

        rows = st.session_state.get("cdc_rows")
        loaded_id = st.session_state.get("cdc_loaded_id")
        if rows is not None and loaded_id == source.dataset_id:
            frame = pd.DataFrame(rows)
            st.success(
                f"Loaded {len(frame):,} row(s) from "
                f"{st.session_state.get('cdc_loaded_title', source.title)}."
            )
            if frame.empty:
                st.info("The query returned no rows.")
            else:
                st.dataframe(frame, use_container_width=True, hide_index=True)
                st.download_button(
                    "Download current preview as CSV",
                    data=frame.to_csv(index=False).encode("utf-8"),
                    file_name=f"cdc-{source.dataset_id}-preview.csv",
                    mime="text/csv",
                    key="cdc_preview_download",
                )

        with st.expander("Registered SODA2 sources"):
            st.dataframe(
                pd.DataFrame(
                    [
                        {
                            "Dataset": item.dataset_id,
                            "Title": item.title,
                            "Category": item.category,
                            "Audience": item.audience,
                        }
                        for item in CDC_SODA2_SOURCES
                    ]
                ),
                use_container_width=True,
                hide_index=True,
            )

    with schedule_tab:
        st.caption("Official CDC adult schedule content loaded through the CDC Content Services API.")
        schedule_label = st.selectbox(
            "Schedule",
            [f"{item.title} · {item.media_id}" for item in CDC_CONTENT_SOURCES],
            key="cdc_schedule_source",
        )
        schedule_source = next(
            item
            for item in CDC_CONTENT_SOURCES
            if schedule_label.endswith(str(item.media_id))
        )
        st.caption(schedule_source.description)

        if st.button("Load official CDC schedule", type="primary", key="load_cdc_schedule"):
            try:
                st.session_state["cdc_schedule_html"] = fetch_content_api_html(
                    schedule_source.media_id
                )
                st.session_state["cdc_schedule_id"] = schedule_source.media_id
            except Exception as exc:
                st.session_state.pop("cdc_schedule_html", None)
                st.error(f"CDC Content Services request failed: {exc}")

        schedule_html = st.session_state.get("cdc_schedule_html")
        if (
            schedule_html
            and st.session_state.get("cdc_schedule_id") == schedule_source.media_id
        ):
            components.html(schedule_html, height=950, scrolling=True)

    with clinical_tab:
        st.caption(
            "CDC implementation data and clinical guidance for vaccine evaluation, "
            "dose validation, occupational risk, and travel documentation."
        )

        st.markdown("#### CDSi — computable immunization logic")
        st.caption(
            "CDC CDSi provides implementation-neutral supporting data and validation test cases "
            "for immunization evaluation and forecasting."
        )
        for item in CDSI_DOWNLOADS:
            left, right = st.columns([3, 1])
            with left:
                st.markdown(f"**{item.title}**")
                st.caption(item.description)
            with right:
                st.link_button("Open / Download", item.url, key=f"cdsi_{item.key}")

        st.markdown("#### Clinical guidance library")
        clinical_categories = ["All"] + sorted(
            {item.category for item in CDC_CLINICAL_SOURCES}
        )
        clinical_category = st.selectbox(
            "Guidance category",
            clinical_categories,
            key="cdc_clinical_category",
        )
        clinical_sources = [
            item
            for item in CDC_CLINICAL_SOURCES
            if clinical_category == "All" or item.category == clinical_category
        ]
        clinical_label = st.selectbox(
            "CDC clinical source",
            [f"{item.title} · {item.key}" for item in clinical_sources],
            key="cdc_clinical_source",
        )
        clinical_source = next(
            item for item in clinical_sources if clinical_label.endswith(item.key)
        )
        st.caption(clinical_source.description)
        st.link_button("Open official CDC guidance", clinical_source.url)

        if st.button(
            "Load guidance tables",
            type="primary",
            key="load_cdc_clinical_tables",
        ):
            try:
                tables = fetch_html_tables(clinical_source.url)
                st.session_state["cdc_clinical_tables"] = tables
                st.session_state["cdc_clinical_key"] = clinical_source.key
            except ValueError:
                st.session_state["cdc_clinical_tables"] = []
                st.session_state["cdc_clinical_key"] = clinical_source.key
                st.info(
                    "This CDC guidance is primarily narrative rather than tabular. "
                    "Use the official CDC guidance link above."
                )
            except Exception as exc:
                st.session_state.pop("cdc_clinical_tables", None)
                st.error(f"CDC clinical-guidance request failed: {exc}")

        clinical_tables = st.session_state.get("cdc_clinical_tables")
        if (
            clinical_tables
            and st.session_state.get("cdc_clinical_key") == clinical_source.key
        ):
            clinical_table_choice = st.selectbox(
                "Guidance table",
                [
                    f"Table {index + 1} · {len(frame):,} rows · {len(frame.columns)} columns"
                    for index, frame in enumerate(clinical_tables)
                ],
                key="cdc_clinical_table_choice",
            )
            clinical_table_index = int(clinical_table_choice.split()[1]) - 1
            clinical_frame = clinical_tables[clinical_table_index]
            st.dataframe(clinical_frame, use_container_width=True, hide_index=True)
            st.download_button(
                "Download displayed guidance table as CSV",
                data=clinical_frame.to_csv(index=False).encode("utf-8"),
                file_name=f"cdc-{clinical_source.key}-table-{clinical_table_index + 1}.csv",
                mime="text/csv",
                key="cdc_clinical_download",
            )

    with codes_tab:
        st.caption(
            "Live CDC IIS reference tables for CVX, MVX, products, NDC, CPT, VIS, "
            "vaccine groups, and current respiratory-season codes."
        )
        st.info(
            "CDC currently warns that its Runtime REST/Viewpoint feeds are temporarily behind "
            "the latest code-set releases, so this app uses CDC's current published reference "
            "tables as the primary source."
        )
        code_label = st.selectbox(
            "Code/reference source",
            [f"{item.title} · {item.key}" for item in CDC_IIS_TABLE_SOURCES],
            key="cdc_code_source",
        )
        code_source = next(
            item for item in CDC_IIS_TABLE_SOURCES if code_label.endswith(item.key)
        )
        st.caption(code_source.description)

        links_col1, links_col2 = st.columns(2)
        with links_col1:
            st.link_button("CDC Runtime REST", IIS_RUNTIME_REST_URL)
        with links_col2:
            st.link_button("CDC code-set access options", IIS_ACCESS_OPTIONS_URL)

        if st.button("Load CDC code tables", type="primary", key="load_cdc_codes"):
            try:
                tables = fetch_html_tables(code_source.url)
                st.session_state["cdc_code_tables"] = tables
                st.session_state["cdc_code_key"] = code_source.key
            except Exception as exc:
                st.session_state.pop("cdc_code_tables", None)
                st.error(f"CDC code-table request failed: {exc}")

        code_tables = st.session_state.get("cdc_code_tables")
        if code_tables and st.session_state.get("cdc_code_key") == code_source.key:
            table_choice = st.selectbox(
                "Table",
                [
                    f"Table {index + 1} · {len(frame):,} rows · {len(frame.columns)} columns"
                    for index, frame in enumerate(code_tables)
                ],
                key="cdc_code_table_choice",
            )
            table_index = int(table_choice.split()[1]) - 1
            code_frame = code_tables[table_index]
            st.dataframe(code_frame, use_container_width=True, hide_index=True)
            st.download_button(
                "Download displayed code table as CSV",
                data=code_frame.to_csv(index=False).encode("utf-8"),
                file_name=f"cdc-{code_source.key}-table-{table_index + 1}.csv",
                mime="text/csv",
                key="cdc_code_download",
            )

    with travel_tab:
        st.caption(
            "Destination-specific CDC Travelers' Health recommendations, including routine "
            "vaccines, travel vaccines, yellow fever recommendations, and country entry requirements."
        )
        st.link_button("CDC Yellow Book country guidance", YELLOW_BOOK_COUNTRY_URL)

        if st.button("Load CDC destination list", type="primary", key="load_cdc_destinations"):
            try:
                st.session_state["cdc_destinations"] = fetch_travel_destinations()
            except Exception as exc:
                st.session_state.pop("cdc_destinations", None)
                st.error(f"CDC destination-list request failed: {exc}")

        destinations = st.session_state.get("cdc_destinations")
        if destinations:
            destination_name = st.selectbox(
                "Destination",
                list(destinations.keys()),
                key="cdc_destination",
            )
            destination_url = destinations[destination_name]
            st.link_button("Open CDC destination page", destination_url)

            if st.button(
                "Load vaccine recommendations",
                type="primary",
                key="load_cdc_travel_vaccines",
            ):
                try:
                    travel_frame = fetch_travel_vaccine_table(destination_url)
                    st.session_state["cdc_travel_frame"] = travel_frame
                    st.session_state["cdc_travel_name"] = destination_name
                except Exception as exc:
                    st.session_state.pop("cdc_travel_frame", None)
                    st.error(f"CDC travel-vaccine request failed: {exc}")

            travel_frame = st.session_state.get("cdc_travel_frame")
            if (
                travel_frame is not None
                and st.session_state.get("cdc_travel_name") == destination_name
            ):
                st.dataframe(travel_frame, use_container_width=True, hide_index=True)
                st.download_button(
                    "Download destination vaccine recommendations as CSV",
                    data=travel_frame.to_csv(index=False).encode("utf-8"),
                    file_name=f"cdc-travel-{re.sub(r'[^A-Za-z0-9]+', '-', destination_name).strip('-').lower()}.csv",
                    mime="text/csv",
                    key="cdc_travel_download",
                )

    with notices_tab:
        st.caption(
            "Current CDC Travel Health Notices from the official Travelers' Health RSS feed."
        )
        links_left, links_right = st.columns(2)
        with links_left:
            st.link_button("Open Travel Health Notices", TRAVEL_HEALTH_NOTICES_URL)
        with links_right:
            st.link_button("Open CDC notices RSS", TRAVEL_HEALTH_NOTICES_RSS)

        if st.button(
            "Refresh Travel Health Notices",
            type="primary",
            key="load_cdc_travel_notices",
        ):
            try:
                st.session_state["cdc_travel_notices"] = fetch_travel_health_notices()
            except Exception as exc:
                st.session_state.pop("cdc_travel_notices", None)
                st.error(f"CDC Travel Health Notice request failed: {exc}")

        notices = st.session_state.get("cdc_travel_notices")
        if notices:
            notices_frame = pd.DataFrame(notices)
            notice_search = st.text_input(
                "Filter notices by country, disease, or level",
                key="cdc_notice_search",
            ).strip()
            if notice_search:
                mask = notices_frame.astype(str).apply(
                    lambda col: col.str.contains(
                        notice_search, case=False, na=False, regex=False
                    )
                ).any(axis=1)
                notices_frame = notices_frame.loc[mask]

            display_columns = ["level", "title", "published", "link"]
            st.dataframe(
                notices_frame[display_columns],
                use_container_width=True,
                hide_index=True,
                column_config={
                    "link": st.column_config.LinkColumn("CDC notice"),
                },
            )
            st.download_button(
                "Download current notices as CSV",
                data=notices_frame.to_csv(index=False).encode("utf-8"),
                file_name="cdc-travel-health-notices.csv",
                mime="text/csv",
                key="cdc_notices_download",
            )

    with imports_tab:
        st.caption(
            "Downloadable NHIS sources are intentionally parked here for the separate Data Importer workflow."
        )
        st.dataframe(
            pd.DataFrame(PARKED_IMPORT_SOURCES),
            use_container_width=True,
            hide_index=True,
        )


def main() -> None:
    import streamlit as st

    from ui_experience import apply_luminous_ui, render_landing_page

    st.set_page_config(
        page_title="Vaccine Prescription Generator",
        page_icon="💉",
        layout="wide",
        initial_sidebar_state="collapsed",
    )
    apply_luminous_ui()

    if st.query_params.get("view") != "app":
        render_landing_page()

    st.markdown(
        """
        <div class="hero">
          <h1>Vaccine Prescription Generator</h1>
          <p>Generate vaccine prescriptions from saved Word forms in the Neon library. Use <code>{{column_name}}</code> tokens, then generate one document per spreadsheet row.</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    if st.button("Return to landing page"):
        st.query_params.clear()
        st.rerun()

    tab_generate, tab_forms, tab_cdc = st.tabs(
        ["Generate Prescriptions", "Prescription Forms", "CDC Data"]
    )
    with tab_generate:
        render_generate_tab()
    with tab_forms:
        render_configure_tab()
    with tab_cdc:
        render_cdc_data_tab()


if __name__ == "__main__":
    main()
